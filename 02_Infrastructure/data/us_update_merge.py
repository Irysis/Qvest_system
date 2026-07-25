# -*- coding: utf-8 -*-
#==============================================================================
# us_update_merge.py — Universe_Support_update.xlsx 증분 merge (W3, 2026-07-25)
#
# incremental_universe_support() (incremental_update_file.R)의 실행 엔진.
# D1 우회 스크립트(d1_us_merge.py, 2026-07-25)의 로직을 상설 이식 + 결함 수리:
#   [구 결함] parse_universe_support(force=FALSE) 호출 → 캐시 존재 시 전 시트
#             skip(no-op) / force=TRUE 면 update xlsx(스냅샷 4개)만으로 패널
#             통째 덮어쓰기 → 1990~ 역사 소실.
#   [수리 ①] cache-hit 판정 = 데이터 기준: update 시트 max Date > 패널 max Date
#             일 때만 merge 진행 (파일 존재 여부 아님).
#   [수리 ②] merge = 기존 패널에서 update 스냅샷 날짜만 제거 후 append
#             (겹침 날짜 교체, 역사 보존 — clobber 금지).
#   [수리 ③] 컬럼명 정합: 디스크 시맨틱 컬럼(us_k200의 'K200' 등, D1 확립)
#             기준. 레거시 'Value' 패널을 만나면 시맨틱으로 정규화 후 merge.
#   [수리 ④] 쓰기 = temp-rename (부분 쓰기 방어).
#
# 파싱: openpyxl read_only 스트리밍 (484MB xlsx — 시트 XML bloat 때문에
#       R openxlsx/readxl DOM 파싱 회피. D1 실측 완주 경로).
# QT 레이아웃: sheet row 8 = ticker 코드 행(col B~), row 9+ 중 col A가
#       datetime인 행 = 스냅샷 행 (실측 ~17행, 방어 상한 200행).
#
# Usage:
#   python us_update_merge.py --xlsx <update.xlsx> --cache <cache_dir> [--force]
#   --force: cache-hit(무신규) 시트도 겹침 날짜 재파싱·교체 (역사 보존 불변)
# Exit: 0 = 정상(no-op 포함) / 1 = 1개 이상 시트 오류
#==============================================================================
import sys, io, os, argparse, datetime

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook
import pyarrow as pa
import pyarrow.parquet as pq
import pyarrow.compute as pc

# (sheet_name, cache_name, semantic_col, is_numeric)
# parse_universe_support.R::UNIVERSE_SUPPORT_SHEET_META 와 1:1 동일 순서
SHEETS = [
    ("Sector_Lv1",     "sector_lv1",      "Sector_Lv1",     False),
    ("Sector_Lv2",     "sector_lv2",      "Sector_Lv2",     False),
    ("유동주식비율",     "float",           "Float",          True),
    ("K200",           "k200",            "K200",           True),
    ("KQ150",          "kq150",           "KQ150",          True),
    ("거래정지",         "trading_halt",    "TradingHalt",    True),
    ("관리종목",         "admin_stock",     "AdminStock",     True),
    ("불성실공시법인",    "unfaithful_disc", "UnfaithfulDisc", True),
]

MAX_SNAP_ROWS = 200  # 실측 ~17행 — 트레일링 bloat 행 방어 상한


def parse_sheet(wb, sheet_name):
    """row 8 = ticker 코드 (col B~), 이후 col A가 datetime인 행 = 스냅샷 행."""
    ws = wb[sheet_name]
    tickers = None
    snaps = []  # (date, values list aligned to tickers)
    for i, row in enumerate(ws.iter_rows(min_row=8, values_only=True), start=8):
        if i == 8:
            tickers = list(row[1:])
            while tickers and tickers[-1] in (None, ""):
                tickers.pop()
            continue
        a = row[0] if row else None
        if isinstance(a, datetime.datetime):
            snaps.append((a.date(), list(row[1:len(tickers) + 1])))
        if i > 8 + MAX_SNAP_ROWS:
            break
    if not tickers or len(tickers) != len(set(tickers)):
        raise ValueError(f"{sheet_name}: ticker row invalid/duplicated")
    return tickers, snaps


def merge_one(wb, cache_dir, sheet_name, cache_name, col, is_num, force):
    """1개 시트 파싱 → 패널 merge. return: 상태 문자열."""
    pq_path = os.path.join(cache_dir, f"us_{cache_name}.parquet")
    if sheet_name not in wb.sheetnames:
        print(f"[{cache_name}] WARN: sheet '{sheet_name}' not in workbook — skip")
        return "missing_sheet"
    if not os.path.exists(pq_path):
        print(f"[{cache_name}] WARN: panel missing ({pq_path}) — "
              f"full parse_universe_support() 필요 (update만으로 신규 생성 금지: 역사 부재)")
        return "missing_panel"

    t0 = datetime.datetime.now()
    tickers, snaps = parse_sheet(wb, sheet_name)
    if not snaps:
        print(f"[{cache_name}] no snapshot rows in update sheet — skip")
        return "no_data"

    existing = pq.read_table(pq_path)
    # ③ 컬럼명 정합: 시맨틱 기준. 레거시 'Value' 패널은 시맨틱으로 정규화.
    names = existing.schema.names
    if names == ["Date", "Ticker", "Value"]:
        existing = existing.rename_columns(["Date", "Ticker", col])
        print(f"[{cache_name}] legacy 'Value' column normalized -> '{col}'")
    elif names != ["Date", "Ticker", col]:
        raise ValueError(f"{cache_name}: unexpected panel schema {names}")

    panel_max = pc.max(existing.column("Date")).as_py()
    upd_max = max(d for d, _ in snaps)

    # ① cache-hit 판정 (데이터 기준): 신규 날짜 없으면 no-op (force면 겹침 교체 진행)
    if upd_max <= panel_max and not force:
        print(f"[{cache_name}] cache-hit: update max {upd_max} <= panel max {panel_max} — no-op")
        return "noop"

    # 파싱 → long (QT_to_xts / QT_to_dt_char parity: NA·빈값 drop,
    # numeric 시트는 as.numeric 실패값 drop, char 시트는 ''/'NA' drop)
    new_dates, new_tickers, new_vals = [], [], []
    per_date = {}
    for d, vals in snaps:
        n_kept = 0
        for tk, v in zip(tickers, vals):
            if v is None or v == "":
                continue
            if is_num:
                try:
                    fv = float(v)
                except (TypeError, ValueError):
                    continue
                new_vals.append(fv)
            else:
                sv = str(v).strip()
                if sv == "" or sv == "NA":
                    continue
                new_vals.append(sv)
            new_dates.append(d)
            new_tickers.append(str(tk))
            n_kept += 1
        per_date[d] = n_kept

    upd_dates = sorted(per_date)
    # ② merge: 겹침 날짜만 제거 후 append — 역사 보존
    mask = pc.invert(pc.is_in(existing.column("Date"),
                              value_set=pa.array(upd_dates, pa.date32())))
    kept = existing.filter(mask)
    schema = existing.schema
    val_type = schema.field(col).type
    new_tbl = pa.table(
        {"Date": pa.array(new_dates, pa.date32()),
         "Ticker": pa.array(new_tickers, pa.string()),
         col: pa.array(new_vals, val_type)},
        schema=schema)
    combined = pa.concat_tables([kept, new_tbl]).sort_by(
        [("Date", "ascending"), ("Ticker", "ascending")])

    # ④ temp-rename 쓰기
    tmp = pq_path + ".tmp"
    pq.write_table(combined, tmp)
    os.replace(tmp, pq_path)

    mins = (datetime.datetime.now() - t0).total_seconds() / 60
    print(f"[{cache_name}] rows {existing.num_rows} -> {combined.num_rows} | "
          f"replaced/appended dates: {[(str(d), per_date[d]) for d in upd_dates]} "
          f"({mins:.1f}min)", flush=True)
    return "merged"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--cache", required=True)
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    if not os.path.exists(args.xlsx):
        print(f"[us_update_merge] xlsx not found: {args.xlsx}")
        sys.exit(1)

    wb = load_workbook(args.xlsx, read_only=True, data_only=True)
    n_err = 0
    summary = {}
    try:
        for sheet_name, cache_name, col, is_num in SHEETS:
            try:
                summary[cache_name] = merge_one(
                    wb, args.cache, sheet_name, cache_name, col, is_num, args.force)
            except Exception as e:  # noqa: BLE001 — 시트별 격리, 전체 중단 방지
                n_err += 1
                summary[cache_name] = "ERROR"
                print(f"[{cache_name}] ERROR: {type(e).__name__}: {e}", flush=True)
    finally:
        wb.close()

    print("[us_update_merge] summary: "
          + ", ".join(f"{k}={v}" for k, v in summary.items()))
    sys.exit(1 if n_err else 0)


if __name__ == "__main__":
    main()
