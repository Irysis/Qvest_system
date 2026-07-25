# -*- coding: utf-8 -*-
"""us_update_stream_parse.py — Universe_Support_update.xlsx 스트리밍 파서 (증분 표준 경로)

D1(2026-07-25) 우회 스크립트(d1_us_merge.py)의 파싱부를 표준 경로로 승격한 것.
incremental_update_file.R::incremental_universe_support()가 호출한다 — 단독 실행용 아님.

왜 openpyxl read_only 스트리밍인가:
  update xlsx는 시트 XML이 스타일 잔재로 420~560MB bloat 상태(실측 D1). openxlsx/readxl은
  시트 XML 전체를 메모리에 올려 파싱하므로 시트당 GB급 스파이크 — startRow=8 이후 실데이터가
  ~17행뿐이라 스트리밍 + 조기종료(MAX_SCAN_ROW)로 bloat 구간을 아예 읽지 않는다.

레이아웃 (QT 표준): row 8 = 종목코드 헤더(col B~), row 9~13 = 메타, 이후 col A가 날짜인 행 = 스냅샷.
값 필터 = parse_universe_support.R parse_one_us_sheet()와 parity:
  numeric → float 변환 실패/None/'' drop (QT_to_xts as.numeric NA-drop 동형, 0은 보존)
  character → strip 후 ''/'NA' drop

출력: <out_dir>/us_new_<cache_name>.csv (utf-8, 헤더 = Date,Ticker,<시맨틱 컬럼명>)
  parquet가 아닌 CSV인 이유: QVEST_PY(시스템 python312)에 pyarrow 부재 — openpyxl만으로
  동작하게 해 인터프리터 의존을 최소화. 타입/스키마 확정은 R merge 단계가 담당.
stdout: ASCII-only 요약 + 마지막 줄 "PARSE_OK n_sheets=8" 센티널.
  호출측(R)은 exit 0 && PARSE_OK 둘 다 확인(fail-closed — exit 0 성공 오보 방어).

usage: python us_update_stream_parse.py <update_xlsx> <out_dir>
"""
import sys
import os
import io
import csv
import datetime

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook

# (sheet_name, cache_name, semantic_col, is_numeric)
# parse_universe_support.R::UNIVERSE_SUPPORT_SHEET_META와 1:1 — 변경 시 양쪽 동기 의무
SHEETS = [
    ("Sector_Lv1",      "sector_lv1",      "Sector_Lv1",     False),
    ("Sector_Lv2",      "sector_lv2",      "Sector_Lv2",     False),
    ("유동주식비율",     "float",           "Float",          True),
    ("K200",            "k200",            "K200",           True),
    ("KQ150",           "kq150",           "KQ150",          True),
    ("거래정지",         "trading_halt",    "TradingHalt",    True),
    ("관리종목",         "admin_stock",     "AdminStock",     True),
    ("불성실공시법인",   "unfaithful_disc", "UnfaithfulDisc", True),
]

MAX_SCAN_ROW = 300  # 실측 ~17행 — 방어 상한. 조기 break가 시트 XML bloat 미소비의 핵심


def parse_sheet(wb, sheet_name):
    """row 8 = ticker 헤더, 이후 col A가 날짜(datetime 또는 Excel 직렬수)인 행 = 스냅샷 행."""
    ws = wb[sheet_name]
    tickers = None
    snaps = []  # (date, values aligned to tickers)
    for i, row in enumerate(ws.iter_rows(min_row=8, values_only=True), start=8):
        if i == 8:
            tickers = list(row[1:])
            while tickers and tickers[-1] in (None, ""):
                tickers.pop()
            continue
        a = row[0] if row else None
        d = None
        if isinstance(a, datetime.datetime):
            d = a.date()
        elif isinstance(a, datetime.date):
            d = a
        elif isinstance(a, (int, float)) and not isinstance(a, bool) and 20000 < a < 60000:
            # 날짜 서식 미적용 Excel 직렬수 방어 (col A는 날짜 또는 메타 텍스트만)
            d = (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(a)))
        if d is not None:
            snaps.append((d, list(row[1 : len(tickers) + 1])))
        if i > MAX_SCAN_ROW:
            break
    if not tickers:
        raise ValueError(f"{sheet_name}: ticker header row(8) empty")
    tickers = [str(t) for t in tickers]
    if len(tickers) != len(set(tickers)):
        raise ValueError(f"{sheet_name}: duplicated ticker in header row")
    return tickers, snaps


def main():
    if len(sys.argv) != 3:
        print("usage: us_update_stream_parse.py <update_xlsx> <out_dir>")
        return 2
    xlsx, out_dir = sys.argv[1], sys.argv[2]
    if not os.path.exists(xlsx):
        print(f"ERROR xlsx not found: {xlsx}")
        return 2
    os.makedirs(out_dir, exist_ok=True)

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    n_ok = 0
    try:
        for sheet_name, cache_name, col, is_num in SHEETS:
            t0 = datetime.datetime.now()
            if sheet_name not in wb.sheetnames:
                print(f"ERROR sheet missing in xlsx: {cache_name}")
                return 3
            tickers, snaps = parse_sheet(wb, sheet_name)

            out_csv = os.path.join(out_dir, f"us_new_{cache_name}.csv")
            n_rows = 0
            per_date = []
            with open(out_csv, "w", newline="", encoding="utf-8") as fh:
                w = csv.writer(fh)
                w.writerow(["Date", "Ticker", col])
                for d, vals in snaps:
                    n_kept, ssum = 0, 0.0
                    for tk, v in zip(tickers, vals):
                        if v is None or v == "":
                            continue
                        if is_num:
                            try:
                                fv = float(v)
                            except (TypeError, ValueError):
                                continue  # as.numeric NA-drop parity
                            w.writerow([d.isoformat(), tk, repr(fv)])
                            ssum += fv
                        else:
                            sv = str(v).strip()
                            if sv in ("", "NA"):
                                continue
                            w.writerow([d.isoformat(), tk, sv])
                        n_kept += 1
                        n_rows += 1
                    per_date.append(
                        (d.isoformat(), n_kept, round(ssum, 2) if is_num else None)
                    )
            secs = (datetime.datetime.now() - t0).total_seconds()
            print(
                f"[{cache_name}] rows={n_rows} dates={len(snaps)} "
                f"per_date={per_date} ({secs:.0f}s)",
                flush=True,
            )
            n_ok += 1
    finally:
        wb.close()

    print(f"PARSE_OK n_sheets={n_ok}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
